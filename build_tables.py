"""Build outputs/final_tables.md + outputs/s7_funnel_behaviour_aug_sep.xlsx from the sql/03, 14, 15 CSVs."""
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

BASE = 'A00 All registrations (reached S4) = overall'
H = ['Segment', 'Users', 'Share of S6', 'Sales', 'Conv %', 'Overall conv %', 'x overall',
     'Attended', 'WA %', 'Overall WA %', 'x overall', 'WA % (held)']
PL = {'Aug+Sep': 'Aug + Sep', '2026-08': 'August', '2026-09': 'September'}

wb = Workbook(); wb.remove(wb.active)
FILL = PatternFill('solid', fgColor='1F4E78'); HF = Font(bold=True, color='FFFFFF')

def sheet(name, blocks):
    ws = wb.create_sheet(name); r = 1
    for title, head, rows in blocks:
        ws.cell(r, 1, title).font = Font(bold=True, size=12); r += 1
        for j, h in enumerate(head, 1):
            c = ws.cell(r, j, h); c.fill = FILL; c.font = HF; c.alignment = Alignment(wrap_text=True)
        r += 1
        for row in rows:
            for j, v in enumerate(row, 1): ws.cell(r, j, v)
            r += 1
        r += 2
    ws.column_dimensions['A'].width = 42
    for col in 'BCDEFGHIJKLMNOPQR': ws.column_dimensions[col].width = 13

def load(path, keyf):
    D = {}
    for r in csv.DictReader(open(path)):
        D[(keyf(r), r['segment'])] = [int(r[k]) if r[k] else None for k in ('users', 'sales', 'attended', 'reg_held', 'att_held')]
    return D

def build(D, key):
    segs = sorted(s for (k, s) in D if k == key and not s.startswith('Z'))
    b = D[(key, BASE)]; s6 = D[(key, 'A01 Reached S6')][0]
    bc, bw = b[1] / b[0], b[2] / b[0]
    out = []
    for s in segs:
        n, sa, at, rh, ah = D[(key, s)]
        share = '' if s[0] == 'D' or s == BASE else f'{100*n/s6:.1f}%'
        name = s[4:]
        if s[:3] in ('B03', 'B04', 'B05'): name = '  - ' + name
        out.append([name, n, share, sa, f'{100*sa/n:.2f}%', f'{100*bc:.2f}%', f'{sa/n/bc:.2f}x', at,
                    f'{100*at/n:.1f}%', f'{100*bw:.1f}%', f'{at/n/bw:.2f}x', f'{100*ah/rh:.1f}%' if rh else 'n/a'])
    return out

def recon(D, keys):
    return [[PL[k[0] if isinstance(k, tuple) else k], D[(k, 'Z01 Reached S4, before lead match')][0], D[(k, BASE)][0],
             f"{100*D[(k, BASE)][0]/D[(k, 'Z01 Reached S4, before lead match')][0]:.1f}%"] for k in keys]
RH = ['Period', 'S4 users (clickstream)', 'S4 users with a lead (table base)', 'Matched']

md = []
f = [r for r in csv.DictReader(open('outputs/03_funnel_mom.csv')) if r['month'] >= '2026-08']
FH = ['Month', 'S0', 'S1', 'S2', 'S3', 'S4', 'S5', 'S6', 'S7', 'S2/S1', 'S3/S2', 'S4/S3', 'S4/S1', 'S5/S4', 'S6/S5', 'S7/S6', 'S7/S1', 'S7/S4']
frows = []
for r in f:
    g = lambda k: int(r[k]); pc = lambda a, b: f'{100*g(a)/g(b):.2f}%'
    frows.append([{'2026-08': 'Aug', '2026-09': 'Sep (to 28th)'}[r['month']]] + [g(k) for k in ['s0', 's1', 's2', 's3', 's4', 's5', 's6', 's7']] +
                 [pc('s2', 's1'), pc('s3', 's2'), pc('s4', 's3'), pc('s4', 's1'), pc('s5', 's4'), pc('s6', 's5'), pc('s7', 's6'), pc('s7', 's1'), pc('s7', 's4')])
sheet('1 Funnel MoM', [('Overall funnel, US, unique users per step per month', FH, frows)])
md.append(('## 1. Overall funnel, month on month', FH, frows))

D = load('outputs/14_conv_wa_unique_users_aug_sep.csv', lambda r: r['period'])
sheet('2 Behaviour overall', [(f'Overall - {PL[p]}', H, build(D, p)) for p in PL] + [('Reconciliation to funnel S4', RH, recon(D, list(PL)))])
md.append(('## 2. User behaviour vs conversion & WA, overall (Aug + Sep, unique users)', H, build(D, 'Aug+Sep')))
md.append(('### Reconciliation: S4 users vs table base', RH, recon(D, list(PL))))

D2 = load('outputs/15_conv_wa_paid_nonpaid_aug_sep.csv', lambda r: (r['period'], r['grp']))
for g, tag in [('Paid', '3a'), ('Non-paid', '3b')]:
    keys = [(p, g) for p in PL]
    sheet(f'3 Behaviour {g}', [(f'{g} - {PL[p]}', H, build(D2, (p, g))) for p in PL] + [('Reconciliation', RH, recon(D2, keys))])
    md.append((f'## {tag}. {g} (Aug + Sep, unique users)', H, build(D2, ('Aug+Sep', g))))

wb.save('outputs/s7_funnel_behaviour_aug_sep.xlsx')
with open('outputs/final_tables.md', 'w', encoding='utf-8') as o:
    for t, h, rows in md:
        o.write(f'\n{t}\n\n| ' + ' | '.join(h) + ' |\n|' + '---|' * len(h) + '\n')
        for r in rows: o.write('| ' + ' | '.join(f'{v:,}' if isinstance(v, int) else str(v) for v in r) + ' |\n')
print(open('outputs/final_tables.md', encoding='utf-8').read())
